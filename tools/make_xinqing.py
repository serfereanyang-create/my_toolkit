# -*- coding: utf-8 -*-
"""生成《心情》：正文约1000字，含真正的 Word 目录域（TOC field）。"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"C:\Users\34516\Desktop\心情.docx"

TITLE = "心 情"

SECTIONS = [
    ("一、心情是房间里的光", [
        "最近常想起一个说法：心情不是天气，它更像房间里的光。同一个屋子，早晨是斜的，下午是平的，阴天就整间发灰；屋子没变，你看见的东西却不一样了。窗还是那扇窗，桌子还是那张桌子，可你坐在里面的时候，会觉得这地方今天格外空，或者格外暖。",
        "光是不出声的，心情也是。它不敲门，不通报，推门就进来了。你正低头做事，忽然发现手里的句子读了三遍也不进脑子，才知道它已经在屋里坐了一会儿。有时候它来得有影子，比如一句没头没尾的话，一个忽然安静下来的群聊；有时候它连影子都没有，只是从早上睁眼起，天就比昨天低一些。",
        "所以我不太相信“我今天心情很好”这种话能撑一整天。光会移，人会变，下午的你未必还认同上午的判断。这不是不坚定，这是活着本来的样子。",
    ]),
    ("二、不必为心情找理由", [
        "我以前总觉得心情需要被解释。高兴要说出理由，难过要找到源头，仿佛不给它一个名分，它就不算数、不合法。后来发现这样太累。有些日子就是没来由地沉，像水杯底下积了一层细灰，晃一晃才看得见；也有些日子莫名其妙轻快，走在路上会为一句不相干的广告词笑出声。",
        "它们不需要理由，只需要被允许存在。人的心不是账本，不必每一笔都记清来处。你追问得越紧，它退得越远，最后你手里只剩一个空问题：我到底为什么不高兴？这个问题本身，往往比不高兴更消耗人。",
        "我后来学会的是把问题换成陈述：今天就是低。不解释，不追责，不给自己写检讨。承认之后，那个模糊的、压在胸口的东西反而有了边界，像一个原本铺满整间屋子的影子，被灯一照，缩成了墙角的一块。",
    ]),
    ("三、最难的是承认", [
        "真正难的倒不是坏心情，是不肯承认自己有心情。我们太习惯把它当成需要处理的故障：赶紧好起来，别影响工作，别让人看出来。于是它被按下去，然后换个地方冒出来——变成失眠，变成对无关的人发火，变成对着屏幕发呆半小时却不知道自己想干什么。",
        "它还会伪装。有的时候是忽然对一切都不耐烦，有的时候是拼命找事做，把日程排满，让自己没有空隙去想。忙起来确实不难受，但那是把难受推迟，不是消化。账总要还的，通常在你最没防备的深夜。",
        "承认反而简单：哦，我今天不太行。就这一句，说出口，它就小了一号。像屋里进了只飞虫，你打开灯看清它，它反而没那么可怕了。真正可怕的从来是黑暗里那个不知道多大的东西。",
    ]),
    ("四、与它和平共处", [
        "所以现在我不太跟心情讲道理了。它来，我给它倒杯水；它走，我不追。我学着看它的形状：是一阵能过去的风，还是一层压得很低的云。风就让它过，云就早点睡。能做的事其实很少，但很少也够了。",
        "具体的办法都是笨办法。把窗户打开，让真的空气进来；洗个热水澡；出门走二十分钟，不看手机；把该做的事拆成小得可笑的一步，只做那一步。这些办法治不好什么，但它们能把人从原地挪开一点点，而挪开一点点，往往就是全部的区别。",
        "日子还是一样长，但好像能喘上气了。我不再要求自己每天都是晴天，也不再把阴天当成失败。人这一辈子没法一直晴着，能做的大概只是把屋子收拾得干净些，让不管什么光进来，都照得见一点东西。",
    ]),
    ("五、留一点余地", [
        "写到这里，窗外已经暗了。屋子里的光从斜的变成平的，又变成一片模糊。我并没有因为写完这篇东西就变得更好，只是比早上更清楚自己在哪儿。",
        "这大概就是我想要的全部：不指望心情听话，只希望自己别跟它打架。它来的时候我在，它走的时候我也在。中间这些时间，够我读完一页书，煮一碗面，跟人说一句真心的话。",
        "心情不是要被解决的问题，它是你活着这件事的回声。听见了，点个头，然后继续走。",
    ]),
]

TIP = "更新目录：在目录上右键 → 更新域 → 更新整个目录（或按 Ctrl+A 后按 F9）。"


def set_cn_font(run, name="宋体", size=None, bold=False):
    run.font.name = name
    run.font.bold = bold
    if size:
        run.font.size = size
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def add_toc_field(doc):
    """插入真正的 TOC 域，Word 打开后可更新出标准引用目录。"""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "（目录域已插入，右键“更新域”生成本文各级标题与页码）"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr, fld_sep, placeholder, fld_end):
        run._r.append(el)
    set_cn_font(run, "宋体", Pt(10.5))


def enable_update_fields(doc):
    """让 Word 打开文档时自动提示更新域，目录立刻成形。"""
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)


def main():
    doc = Document()

    # 默认样式改成中文论文口径
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)

    for lvl in (1, 2):
        st = doc.styles["Heading %d" % lvl]
        st.font.name = "黑体"
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.bold = True
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    # 标题
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_after = Pt(18)
    set_cn_font(t.add_run(TITLE), "黑体", Pt(22), bold=True)

    # 目录二字
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.paragraph_format.space_after = Pt(10)
    set_cn_font(h.add_run("目　　录"), "黑体", Pt(14), bold=True)

    add_toc_field(doc)

    tip = doc.add_paragraph()
    tip.paragraph_format.space_before = Pt(4)
    tip.paragraph_format.space_after = Pt(18)
    set_cn_font(tip.add_run(TIP), "楷体", Pt(9))

    # 正文
    for idx, (head, paras) in enumerate(SECTIONS):
        hp = doc.add_heading(level=1)
        hp.paragraph_format.space_before = Pt(12)
        hp.paragraph_format.space_after = Pt(6)
        set_cn_font(hp.add_run(head), "黑体", Pt(14), bold=True)
        for txt in paras:
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(24)  # 首行缩进2字符
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(6)
            set_cn_font(p.add_run(txt), "宋体", Pt(12))

    enable_update_fields(doc)
    doc.save(OUT)

    body = "".join("".join(ps) for _, ps in SECTIONS)
    cjk = sum(1 for ch in body if "\u4e00" <= ch <= "\u9fff")
    print("saved:", OUT)
    print("cjk_chars:", cjk)
    print("sections:", len(SECTIONS))


if __name__ == "__main__":
    main()
